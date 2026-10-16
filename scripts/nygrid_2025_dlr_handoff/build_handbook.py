"""Build the colleague handoff from reviewed model evidence and portable APIs.

Run with the Codex bundled Python when authoring the DOCX. Runtime researchers
do not need python-docx: the generated handbook is included in the packages.
"""
from pathlib import Path
from copy import deepcopy
import hashlib
import json
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'output/nygrid_2025_dlr_handoff'
DOCX=ROOT/'NYGRID_2025_DLR_HANDOFF.docx'
TEXT=ROOT/'NYGRID_2025_DLR_HANDOFF.md'
BASELINE='8188cbec1172219a0414881121ad042b405a4044'
TITLE='New York grid model handoff for DLR research'
SOURCES={
 '2019 parent':'NYGRID_2019_DLR_REBUILD.md',
 '2025 inputs':'NYGRID_2025_OPERATING_INPUTS.md',
 'Fleet and dispatch':'NYGRID_2025_BOUNDED_DISPATCH.md',
 'Downstate construction':'NYGRID_2025_DOWNSTATE_REFINEMENT.md',
 'Small bus DC model':'NYGRID_2025_SLIM_DOWNSTATE.md',
 'Moses diagnosis':'NYGRID_2025_MOSES_SOUTH_DIAGNOSIS.md',
 'Ontario controls':'NYGRID_2025_VERIFIED_PAR_CONTROLS.md',
 'Western source audit':'NYGRID_2025_WESTERN_SOURCE_AUDIT.md',
 'Physical flowgates':'NYGRID_2025_PHYSICAL_FLOWGATE_DEFINITIONS.md',
 'Line estimates':'NYGRID_2025_LINE_PARAMETER_ESTIMATION.md',
 'Dover parameters':'NYGRID_2025_DOVER_PARAMETERS_APPLIED.md',
 'P34 qualifications':'NYGRID_2025_P34_QUALIFICATION_AND_DLR_READINESS.md',
 'Engineering AC and DLR':'NYGRID_2025_AC_SYSTEM.md',
 'Profile AC audit':'NYGRID_2025_ACPF_PROFILE_AUDIT.md',
 'Compact AC':'NYGRID_2025_COMPACT_AC.md',
 'Central East diagnosis':'NYGRID_2025_CENTRAL_EAST_DIAGNOSIS.md',
 'Loss revision':'NYGRID_2025_AC_LOSS_REVISION.md',
 'Branch classification':'output/nygrid_2025_ac_system/branches/branch_assignments.csv',
 'Thermal segments':'output/nygrid_2025_ac_system/branches/thermal_segments.csv',
 'Conductor catalog':'output/nygrid_2025_ac_system/engineering/thermal_conductors.csv',
 'Current physical exceptions':'output/nygrid_2025_acpf_profile_audit/structure/structure_audit.json',
 'Retained bus ledger':'output/nygrid_2025_compact_ac/independent/bus_retention_ledger.json',
 'Final annual loss receipt':'output/nygrid_2025_ac_loss_revision/validation/annual/released_equivalent_boundary_consistent/annual_loss_policy_receipt.json',
 'Final annual comparisons':'output/nygrid_2025_ac_loss_revision/comparison/combined/public_interface_metrics.csv',
 'Final hour masks':'output/nygrid_2025_ac_loss_revision/comparison/combined/comparison_cohort.csv',
 'Final independent physics':'output/nygrid_2025_ac_loss_revision/impedance/released_equivalent_boundary_consistent_annual_arrays_independent_review.json',
 'Earlier project hierarchy':'PROJECT_HANDOFF.md',
 'Phase 1C implementation':'ELECTRICAL_MODEL_IMPLEMENTATION.md',
 'Package A foundation':'output/ny_only_package_a/NY_ONLY_FOUNDATION_RESULTS.md',
 'Package B candidate':'NY_ONLY_PACKAGE_B_IMPLEMENTATION.md',
 '51 bus testbed':'COMPACT_NPCC_NY_TESTBED.md',
 '66 bus calibration':'COMPACT_NY_2025_DLR_METHODOLOGY.md',
 'Independent generation study':'COMPACT_NY_GENERATION_RECONSTRUCTION.md',
 '71 bus partial SPC':'COMPACT_NY_PARTIAL_SPC_BASELINE.md',
 '2019 measurement repair':'NYGRID_2019_INTERFACE_VERIFICATION.md',
 '2019 comparator':'NYGRID_COMPACT_2019_COMPARISON.md',
}


def t(value):
    el=OxmlElement('m:r');text=OxmlElement('m:t');text.text=value;el.append(text);return el


def sub(a,b):
    el=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(t(a));s=OxmlElement('m:sub');s.append(t(b));el.extend([e,s]);return el


def sup(a,b):
    el=OxmlElement('m:sSup');e=OxmlElement('m:e');e.append(a if not isinstance(a,str) else t(a));s=OxmlElement('m:sup');s.append(t(b));el.extend([e,s]);return el


def frac(a,b):
    el=OxmlElement('m:f')
    for key,value in [('num',a),('den',b)]:
        e=OxmlElement('m:'+key)
        for item in value:e.append(item if not isinstance(item,str) else t(item))
        el.append(e)
    return el


def math(doc,items):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    mp=OxmlElement('m:oMathPara');om=OxmlElement('m:oMath')
    for item in items:om.append(item if not isinstance(item,str) else t(item))
    mp.append(om);p._p.append(mp);return p


def hyperlink(p,label,target):
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr');style=OxmlElement('w:rStyle');style.set(qn('w:val'),'Hyperlink');rp.append(style);r.append(rp)
    tx=OxmlElement('w:t');tx.text=label;r.append(tx);link.append(r);p._p.append(link)


def table(doc,headers,rows,widths):
    tab=doc.add_table(rows=1,cols=len(headers));tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
    props=tab._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'5');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    props.append(borders)
    for col,w in zip(tab.columns,widths):col.width=Inches(w)
    for rowidx,values in enumerate([headers]+rows):
        row=tab.rows[0] if rowidx==0 else tab.add_row()
        for j,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');marg.append(e)
            tcPr.append(marg);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'244253' if rowidx==0 else ('F2F5F6' if rowidx%2==0 else 'FFFFFF'));tcPr.append(shade)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1.02
            run=p.add_run(str(value));run.font.size=Pt(9)
            if rowidx==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
        if rowidx==0:
            trPr=row._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');trPr.append(repeat)
        trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return tab


def main():
    OUT.mkdir(exist_ok=True,parents=True)
    summary=json.loads((ROOT/'output/nygrid_2025_ac_loss_revision/summary/revision_summary.json').read_text())
    assert summary['counts']['revised_qualified']==8170
    d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=Inches(.66);s.bottom_margin=Inches(.66);s.left_margin=Inches(.8);s.right_margin=Inches(.8)
    s.header_distance=Inches(.27);s.footer_distance=Inches(.27)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_after=Pt(7)
    d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.line_spacing=1.08
    d.styles['Title'].font.size=Pt(26);d.styles['Title'].font.bold=True
    d.styles['Subtitle'].font.size=Pt(12)
    d.styles['Heading 1'].font.size=Pt(18);d.styles['Heading 2'].font.size=Pt(12);d.styles['Heading 2'].font.bold=True
    for name in ['Heading 1','Heading 2','Heading 3']:d.styles[name].paragraph_format.keep_with_next=True
    code_style=d.styles.add_style('Research Code',1);code_style.font.name='Consolas';code_style.font.size=Pt(8.5)
    code_style.paragraph_format.line_spacing=1;code_style.paragraph_format.space_after=Pt(5)
    code_style.paragraph_format.keep_together=True
    header=s.header.paragraphs[0];header.text='NYgrid 2025  |  DLR research handoff';header.style='Caption';header.runs[0].font.color.rgb=RGBColor(0,0,0)
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('13 October 2026  |  Page ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    d.core_properties.title=TITLE;d.core_properties.subject='Model construction, reproducibility and qualified DLR research use'
    d.core_properties.author='NYgrid research project';d.core_properties.keywords='NYgrid, NPCC, MATPOWER, AC power flow, DLR, research handoff'
    md=['# '+TITLE,'','Prepared 13 October 2026. Model snapshot '+BASELINE+'.','']
    pages=[]
    def para(text):d.add_paragraph(text);md.extend([text,''])
    def heading(text):d.add_paragraph(text,'Heading 2');md.extend(['### '+text,''])
    def page(title):
        if pages:d.add_page_break()
        pages.append(title);d.add_paragraph(title,'Heading 1');md.extend(['## '+title,''])
    def code(text):d.add_paragraph(text,'Research Code');md.extend(['```text',text,'```',''])
    def tab(headers,rows,widths):
        table(d,headers,rows,widths);md.extend(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |'])
        md.extend('| '+' | '.join(str(v).replace('\n','; ') for v in row)+' |' for row in rows);md.append('')
    def refs(*labels):
        p=d.add_paragraph();p.paragraph_format.space_before=Pt(4)
        p.add_run('Evidence  ').bold=True
        for i,label in enumerate(labels):
            if i:p.add_run('  ·  ')
            hyperlink(p,label,SOURCES[label])
        md.extend(['Evidence: '+' · '.join('['+x+']('+SOURCES[x]+')' for x in labels),''])
    def equation(items,plain):math(d,items);md.extend(['```text',plain,'```',''])

    # Deliberate chapter boundaries keep the manual navigable in print and Word.
    d.add_paragraph(TITLE,'Title');d.add_paragraph('Model construction and reproducible experiments','Subtitle')
    para('Prepared for a colleague continuing dynamic line rating research. This package brings together the model builders, public-data reconstructions, electrical and thermal assumptions, saved experiments, and independent validation completed through 13 October 2026. The model snapshot is commit '+BASELINE+'. Packaging utilities and this manual form a new handoff layer; the earlier research evidence is preserved.')
    para('Use the latest 148-bus loss-consistent AC solver for equation and accounting studies, and its recovered original-branch currents for explicitly assumed thermal screening. It is not a validated historical operating case: 8,170 hours satisfy its AC and loss-accounting equations, but none satisfies the joint voltage, reactive and generator-capability gate. The existing bounded AC DLR optimizer belongs to the earlier engineering case and has a separate, limited qualification record.')
    tab(['Research use','Status at handoff'],[
        ['Compact versus detailed AC mathematics','Validated on the stated cases and annual saved-array checks.'],
        ['Public profile and interface comparison','Reproducible with explicit scale, masks, terminal conventions and incomplete physical flowgate mapping.'],
        ['Synthetic conductor and weather screening','Usable on approved same-voltage overhead candidates, with declared assumptions and operating-point flags.'],
        ['Historical asset DLR or operational limits','Not established. Actual conductor, critical-span weather, equipment limits and feasible operating states need further evidence.']], [2.25,4.65])
    heading('Read this first')
    para('Start with the setup and smoke test. Then choose one model family deliberately, inspect the bus and branch ledgers, and keep electrical qualification, physical asset evidence and thermal qualification separate. Never assign conductor physics to a reduced admittance coupling or infer an installed ampacity from a zero MATPOWER rating.')
    refs('Loss revision','Engineering AC and DLR','Compact AC')
    pages.append('Purpose and status')

    page('Start and verify the package')
    para('Extract the archive into a writable directory. The package root contains System Matpower Format, scripts, output and vendor folders. Paths below are relative to that root, so no original user directory is needed. Use a licensed MATLAB installation. The setup helper uses the bundled MATPOWER 8.1 code by default; an explicit compatible external MATPOWER directory may be supplied.')
    code("cd('C:\\research\\nygrid_dlr_handoff');  % your extracted directory\naddpath(fullfile(pwd,'scripts','nygrid_2025_dlr_handoff','runtime'));\nenv = setup_nygrid_2025_handoff(pwd);\ncheck = smoke_nygrid_2025_handoff(fullfile(pwd,'runtime_results'));")
    para('The smoke test checks representative topology, loss accounting, compact/current recovery and a declared thermal screening scenario. Read its saved receipt and any failed flags before using subsequent outputs. A successful smoke test establishes the packaged workflow on that machine; it does not turn the synthetic model into an operationally qualified grid.')
    heading('Run the current AC model')
    code("[r, audit, full, allocation] = ...\n    run_nygrid_2025_loss_consistent_acpf(8508);\nassert(r.success && audit.loss_allocation_gate);\nassert(audit.reduction_and_equation_gate);\nIseries_A = r.recovered_line_current_A(:,3);\nphysical_ok = audit.VQ_and_generator_assumption_gate;")
    para('Hour 8508 is a tested post-Dover example. For this point, the accounting and equation gates pass while the physical operating gate remains false. r.bus contains 148 retained buses; full.bus and full.branch contain the detailed network for that dated state. Current columns are from terminal, to terminal and series amperes, on the row axis of full.branch. Preserve the associated branch keys and presence/status masks.')
    heading('Offline integrity')
    para('Use the package verifier before running modified experiments. It checks the package file manifest and hashes, independently of Git. Original evidence receipts may contain their historical machine paths and commits; these are provenance records, not installation instructions. The handoff includes a mapping of packaged inputs and a separate clean-directory execution receipt.')
    refs('Loss revision','Retained bus ledger')

    page('Choose the correct model family')
    para('Several preserved models have similar names and different scientific contracts. The current NYgrid pathway starts from the pinned Cornell NPCC source and retains all 140 original buses. It is separate from the older S7 to S13 construction and S14 planning pathway, and from the earlier 51, 66 and 71 bus compact calibration studies. A generic older latest-model file must not select the research case implicitly.')
    tab(['Family','Structure and purpose'],[
        ['2019 released and full reconstruction','Released 57-bus DC model; separate 140-bus full parent with 225 AC records, 271 generators and four controllable-link records.'],
        ['2025 source-conditioned DC','140 original buses and 1,464 generator rows, followed by 191/192-bus Downstate detail and a 148-bus DC reduction.'],
        ['2025 engineering AC','266 buses / 485 branch records before full-service Dover; 267 / 488 afterward. Complete synthetic AC parameters, numerical sections and explicit limits.'],
        ['148-bus compact AC','Complex-admittance elimination of passive buses. Recover the detailed network; do not reinterpret reduced couplings as lines.'],
        ['Loss-consistent AC revision','Same 148 retained voltages with released parent-224 equivalent and jointly solved boundary-consistent loss accounting. Preferred current PF entry point.'],
        ['Earlier bounded AC DLR','Legacy engineering case with steady weather ampacity and bounded ACOPF. Four of eight archived weather snapshots qualify.'],
        ['Earlier S7 S13 S14 and compact studies','Preserved provenance, calibration and alternative-method evidence. Their bus counts, scales and validation do not transfer to the current family.']], [1.65,5.25])
    para('A controllable MATPOWER DC record is a modeling abstraction, not necessarily a separately identified physical HVDC asset. The current six records include external-device representations and Ontario PAR flow substitutions. Select models by public function and explicit variant, not by a MAT filename containing final or latest.')
    refs('2019 parent','2025 inputs','Small bus DC model','Compact AC')

    page('Network coverage and dated topology')
    para('The 148 retained buses comprise all 140 original NPCC buses plus eight 345-kV transmission stations. The original set contains 46 New York buses and 94 external buses; retaining the neighboring networks preserves their electrical distribution effects. Added stations are named ports, not extra fitted loads or generator capacity.')
    tab(['Station','Bus ID','Purpose'],[
        ['East Fishkill','10858','Separates Pleasant Valley and Wood Street circuit sections.'],['Wood Street','10902','Preserves the junction toward Pleasant Valley, East Fishkill and Millwood.'],
        ['Ladentown','10774','Keeps Ramapo and Buchanan corridor sections distinct.'],['Dunwoodie','10593','Westchester receiving port.'],
        ['East Garden City','10589','Long Island receiving port.'],['Gowanus','11129','Brooklyn receiving port.'],['Rainey','11130','Queens receiving port.'],['Cricket Valley','20001','F83 and F84 junction toward New England.']], [1.5,.7,4.7])
    para('The detailed Downstate model originally added 51 terminals. Forty-three of those are passive under the frozen injection placement and are eliminated in the compact model. AC engineering adds 75 zero-injection numerical section buses; they represent line discretization, not new substations. Post-Dover adds one passive regulated port. These yield 118 or 119 eliminated voltages.')
    tab(['State boundary','Interpretation'],[
        ['Before 21 May 2025','Pre-May construction state.'],['21 May 2025 at local midnight','Gowanus–Greenwood RCC activation under the conservative next-midnight rule.'],
        ['30 May 2025 at local midnight','Goethals–Fox Hills RCC activation under the same rule.'],['19 December 2025 at 23 EST','Dover full-service boundary, equal to 20 December at 04 UTC.']], [2.15,4.75])
    para('Use the stored UTC clock and topology-state index. The precise modeled boundary can be an accounting convention derived from day-level commissioning evidence. It does not certify the exact switching instant, every 2025 outage, or a complete year-end asset inventory.')
    refs('Downstate construction','Small bus DC model','Dover parameters')

    page('Operating data and scale')
    para('The current model uses a fixed annual power scale gamma = 0.30012075175285535. It is 9,561 MW divided by the maximum simultaneous 2025 NYISO zonal actual-load sum, 31,857.1773 MW. The numerator is a declared NPCC snapshot anchor, not proof of an observed peak. Public power and model power are retained separately.')
    equation([sub('P','model'),' = γ ',sub('P','public')], 'P_model = gamma * P_public')
    para('Divide aggregate model MW by gamma when comparing with public MW. Do not divide recovered AC amperes by gamma: those currents belong to the stated synthetic network and 100-MVA electrical base. Older DC physical-exposure proxies used a different, explicit public-power lift. Their conversion formula must not be copied into the current AC thermal workflow.')
    para('NYISO actual zonal load is a generation-plus-net-interchange boundary balance. Its 2025 definition is verified in the archived Market Participants User’s Guide. The latest loss revision partitions that balance into modeled delivered bus demand plus assigned losses. This must not be described as measured end-use demand. Static within-zone weights and the equal split between buses 81 and 82 in zone J are model assumptions.')
    tab(['Input','Source and qualification'],[
        ['NY load and fuel mix','NYISO P58 actual load and statewide P63 fuel totals; within-zone spatial allocation is assumed.'],
        ['Native generation','EPA gross-output evidence, EIA unit identities/capability, nuclear status and other bounded source priors. Not observed hourly net unit dispatch.'],
        ['External demand','Historical external bus proportions shaped by NY load; not observed 2025 external-area demand.'],
        ['Interchange and device controls','Frozen regional actual exchanges and separate external transfer channels, with raw sign/unit/time qualifications retained.'],
        ['Reactive data','NY load power factor and generator/converter capability policies are assumptions, not a measured reactive history.']], [1.35,5.55])
    para('The annual axis has 8,760 UTC-referenced positions spanning the local 2025 year and its daylight-saving transitions. The latest source mask qualifies 8,651 hours and excludes 109. Never replace excluded values by zero, forward-fill across those gaps, or use a failed solver iterate as a valid observation. Earlier stages used 8,663 hours; compare errors only on a named common cohort.')
    refs('2025 inputs','Fleet and dispatch','Western source audit','Final hour masks')

    page('Fleet and source corrections')
    para('The fleet register preserves 2,245 unique identities and separates identity, eligibility, electrical landing and hourly capability. The runnable case has 1,436 native rows plus 27 external regional equivalents and one HQ equivalent, totaling 1,464 generator rows. Capacity is counted once per admitted identity; many-to-many EPA reporting relationships do not create extra capacity.')
    tab(['Modification','Reason and boundary'],[
        ['Indian Point retirement','Units 2 and 3 have no 2025 capability and cannot become the reference source. Retired identities remain in the inventory.'],
        ['NYIS membership screen','Native eligibility follows reported balancing-authority scope and corroborating source evidence, not New York State location alone.'],
        ['Bayonne and Linden mapping','Corroborated NYISO resources outside NY State are admitted once; Linden GTG6 and unresolved records are excluded. Native generation is distinct from VFT transfers.'],
        ['Cricket Valley and South Fork','Use unique inventory records and existing statewide fuel totals; avoid adding project output twice. Capability and landing uncertainties remain explicit.'],
        ['Hydro zone and generator landing repairs','Confirmed source corrections were applied in separate revisions and traced in source/repair ledgers. Internal interface scores did not select locations.'],
        ['All-boundary actual interchange','Regional physical-exchange reconstruction replaced mixed scheduled/actual accounting, with a common 8,651-hour comparison cohort.']], [2.05,4.85])
    para('Native hourly dispatch is capacity bounded and independently source conditioned. It is not a unit-commitment solution. Minimum stable operation, ramps, reserves, start costs, synchronized status and storage state of charge have not been established. The latest annual PF holds these native active-power schedules fixed while changing the loss-accounting contract.')
    para('The generator row index, electrical bus ID and reporting-unit identity are different axes. Use the saved corrected generator-to-bus mapping and explicit row IDs when aggregating. Reconstructing bus generation from an earlier raw mapping would undo the confirmed repairs.')
    refs('Fleet and dispatch','Moses diagnosis','Western source audit','Profile AC audit')

    page('Physical detail and project parameters')
    para('Downstate construction retains all original NPCC buses and registers predecessor replacements. The detailed model includes 97 exact historical PERFORM circuit rows; they add source detail, not 97 new 2025 lines. Nine predecessor aggregate branch records are retained inactive. Other attachments and project impedances are explicitly assumed. The 30 edges in the smaller DC equivalent are mathematical couplings, not additional conductors.')
    tab(['Change','Implemented treatment'],[
        ['Downstate station detail','Added 51 source or explicitly assumed terminals; retained eight critical ports in the 148-bus representation.'],
        ['F84 Pleasant Valley to Cricket Valley','Design-backed twin Mallard ACSS family, representative geometry and 14.6-mile route; R/X/B are engineering estimates rather than measured impedance.'],
        ['RCC receiving paths','Dated activation and cable/transformer/PAR representations. Staten route corrected to 9.4 miles; cable installation and conductor details remain unresolved.'],
        ['Dover Y17 and 398','Applied line estimates with distinct NY twin Mallard and CT single Bluebird sections; separate parallel PAR banks preserve equipment meaning.'],
        ['Unbuilt or unresolved project detail','NYES overlap, western additions, exact physical gate membership and other incomplete correspondences remain qualified alternatives or omissions.']], [1.8,5.1])
    para('For F84, the applied 100-MVA-base values are R = 0.0008742145129 pu, X approximately 0.0073369903 pu and total B approximately 0.1214494002 pu. Manufacturer temperature points, bundle geometry, representative phase spacing and capacitance simplifications are recorded. The 2,214-A phase rating is a design reference, not a verified terminal-limited facility rating; it is 1,107 A per subconductor for the twin bundle.')
    para('Use the full-precision CSV ledgers when modifying parameters. Preserve total corridor R/X/B when splitting numerical sections, and retain the parent-to-section map. Circuit count and bundle count are separate: parallel circuits change electrical sharing, while subconductors within a phase bundle also change conductor current and geometry.')
    refs('Downstate construction','Line estimates','Dover parameters','Branch classification')

    page('Controlled transfers and interface definitions')
    para('The six prescribed transfer records and their signs are part of the electrical contract. Ontario’s two PAR-flow channels replace the original aggregate branch-211 transfer at runtime, which stays disabled. HQ and Cedars are accounted once. Merchant transfers already included in regional totals must not be added again as independent external supply.')
    para('The Ontario direction is supported by NYISO’s explicit 3 December 2024 effective-date note. Dover uses two separate parallel bank records and a post-cutover bounded phase reconstruction. Published ratings and angle ranges provide limited engineering information; they do not supply actual hourly tap positions, breaker states, factory R/X, full P/Q limits or emergency capability.')
    tab(['Topic','Required interpretation'],[
        ['P34 raw field','The native header says Flow (MWH). Numeric-as-MW interpretation remains conditional; cadence alone does not justify multiplication by 12.'],
        ['P34 timestamps','Forward-held hourly integration and gap limits are declared policies. Native instantaneous/start/end/average semantics are not fully certified.'],
        ['Dover zero postings','Do not infer bypass position, zero total tie flow or duplicate sensors from equal/zero bank values.'],
        ['Seven interface comparisons','Use frozen original-parent branch membership, signs and physical terminal convention after recovery. Do not sum compact matrix edges.'],
        ['Physical flowgate audit','269 printed entries in 21 seasonal definitions are archived. Complete physical model binding is not established for any of the seven gates.']], [1.6,5.3])
    para('The numerical error series uses inherited diagnostic cuts named Dysinger East, West Central, Total East, Moses South, Central East–VC, UPNY ConEd and SPR/DUN South. A cut’s name does not certify membership of the current official NYISO physical flowgate. Source-side, target-side and inherited signed-literal-from conventions are separate diagnostic quantities, not a statistical confidence interval.')
    para('No internal seven-interface observations tune the current compact or loss-consistent AC solver. External reported controls still condition the electrical input. Treat the comparison as source-conditioned reconstruction, not an unconditional forecast.')
    refs('Ontario controls','Dover parameters','P34 qualifications','Physical flowgates')

    page('Engineering AC construction and its limits')
    para('The AC engineering overlay supplies positive resistance/reactance, charging, fixed compensation, finite equipment assumptions and generator capability data on the detailed network. It is an explicit synthetic construction. Passing its arithmetic and matrix checks does not verify the installed asset interpretation.')
    tab(['Reference class','Records','Interpretation'],[
        ['Overhead corridors','186','Catalog/geometry realizations; 128 preserve inherited R/X and 58 use explicit engineering replacements.'],
        ['Cable or urban equivalents','40','Synthetic cable parameters and static current limits; no overhead heat balance.'],
        ['Transformer-compatible paths','38','Finite engineering assumptions, not verified nameplates.'],
        ['Composite voltage-changing equivalents','65','Undecomposed regional/device mixtures; no verified internal conductor inventory.'],
        ['PAR station and accounting links','10','Separate equipment models.'],['Inactive provenance','74','Preserved records, not active assets.']], [2.25,.6,4.05])
    para('These counts describe the earlier 413-record reference classification. The latest revision reclassifies parent 224 as an unresolved network equivalent and excludes it from conductor claims. Additional source conflicts at parents 300, 309, 310 and 386 prevent treating their inherited overhead labels as installed overhead evidence. The handoff screening filter excludes all five.')
    para('Numerical overhead sections are at most 100 km; selected long cable equivalents are also subdivided. The engineering design adds separate fixed reactors at selected original terminals under a declared 75% charging-compensation rule. Original shunts are preserved. Generator Q and S envelopes use assumed capability policies; voltage limits are 0.95–1.05 pu. The cost objective penalizes departure from priors, so its value is not market cost or monetary benefit.')
    para('Historical bounded validation found four of seven electrical snapshots and four of eight weather/current-constrained snapshots qualified. These finite results belong to the earlier engineering case. They are not evidence that the current annual profile replay or the revised loss formulation satisfies all operating limits.')
    refs('Engineering AC and DLR','Branch classification','Current physical exceptions')

    page('Exact compact AC and current recovery')
    para('The compact solver eliminates only passive, zero-prescribed-injection interior buses. All loads, generators, DC terminals and independently controlled voltages remain retained. Fixed shunts are included in the complex admittance matrix. Resistance, charging, off-nominal taps and phase shifts enter before reduction.')
    equation([sub('V','e'),' = −',sup(sub('Y','ee'),'−1'),sub('Y','ek'),sub('V','k')], 'Ve = -(Yee \\ Yek) Vk')
    equation([sub('Y','red'),' = ',sub('Y','kk'),' − ',sub('Y','ke'),sup(sub('Y','ee'),'−1'),sub('Y','ek')], 'Yred = Ykk - Yke * (Yee \\ Yek)')
    equation([sub('V','full'),' = L ',sub('V','k'),'     ',sub('I','f'),' = ',sub('Y','f'),' L ',sub('V','k'),'     ',sub('I','t'),' = ',sub('Y','t'),' L ',sub('V','k')], 'Vfull = L Vk; If = Yf L Vk; It = Yt L Vk')
    para('The retained nonlinear AC equations use this complex matrix directly. A compact matrix coupling is not a physical branch. The saved detailed branch operators recover original from-terminal, to-terminal and series currents after solving. Their row identities, end voltages and base-current conventions remain part of the output contract.')
    para('The original compact annual replay matched all 8,494 qualified detailed solutions: maximum complex-voltage discrepancy 6.70 × 10⁻¹² pu and maximum original-branch current discrepancy 1.48 × 10⁻⁸ A. It reproduced the detailed model’s interface errors rather than improving them. In the loss revision, all 8,170 qualified annual outputs passed independent reconstructed branch/KCL/current checks; 11 qualified representative cases additionally matched detailed coupled and conventional final-allocation PF solutions.')
    heading('When to rebuild')
    para('Rebuild after changing R(T), X, B, a tap, phase shift, shunt, branch status, a source profile, mapping or retained-node contract. Fingerprints intentionally reject stale operators. If adding an injection or independent voltage controller at an eliminated bus, retain it or derive a justified projection; do not discard that device to preserve the count of 148. The loss-consistent wrapper currently asserts its registered 148-bus contract.')
    refs('Compact AC','Retained bus ledger','Final independent physics')

    page('Boundary consistent AC loss accounting')
    para('The earlier annual AC replay kept lossless schedules and loss-inclusive load fixed, leaving Clay to absorb the AC imbalance. The current revision instead solves voltage and the loss allocation simultaneously while holding native active generation, HQ and six controlled transfers at their original schedules.')
    equation([sub('P','delivered,z'),' + ',sub('L','NY,z'),' = ',sub('P','P58,z')], 'Pdelivered,z + LNY,z = PP58,z')
    para('Within-NY corridor losses are split equally between their endpoint zones, or wholly assigned to the common zone. Each external region supplies losses of its internal AC corridors and the AC corridors crossing into New York. Its delivered-load proxy remains fixed, and its generators share the loss increment through frozen bounded participation. This preserves net import at the NY terminals under the current external-area connectivity contract.')
    equation([sub('P','loss,MW'),' = ',sub('S','base,MVA'),' ',sub('R','pu'),' |',sub('I','series,pu'),sup('|','2')], 'Ploss_MW = baseMVA * Rpu * abs(Iseries_pu)^2')
    para('Native real power stays fixed; Q demand scales with modeled delivered P at the inherited power factor. External P limits come from the original finite parent capability boxes, and their assumed Q-circle slices update with P. These are synthetic proxy capabilities. The default PF does not enforce all Q limits or optimize dispatch. A valid electrical/accounting result can therefore fail the physical operating gate.')
    para('The solver uses analytic derivatives of the loss-dependent injections and line search within the positive delivered-load and finite-capability domain. It requires the current zero bus-conductance and zero DC-loss assumptions, and no AC connections between different external regions. Unsupported cases must receive a new explicit allocation contract. Loss clipping, negative served demand and an unlimited reference fallback are not allowed.')
    heading('Gilboa Buchanan equivalent')
    para('Parent 224 restores the released total R/X/B = 0.02/0.0147/0 pu in its two existing numerical sections. It is an added E–G network equivalent, not an identified installed direct line. The paper gives X approximately 0.0146498635 pu; the released 0.0147 differs by 0.342%, not exact four-decimal rounding. No physical R/B derivation was established. Its rating, conductor and R(T) claims are removed, and every other branch parameter and shunt is preserved.')
    refs('Central East diagnosis','Loss revision')

    page('Annual results and interpretation')
    para('The completed loss-consistent year attempts the original 8,651 eligible hours. It qualifies 8,170 and rejects 481, with zero runtime exceptions. All qualified revised hours belong to the old 8,494-success cohort: 324 old successes are lost and none is gained. Failures comprise 465 line-search failures and 16 near-boundary delivered-load failures. Nonconvergence is not proof of global infeasibility.')
    rows=[[r['interface'],f"{r['old_WAPE_percent']:.2f}",f"{r['revised_WAPE_percent']:.2f}"] for r in summary['paired_annual_interfaces']]
    tab(['Interface','Previous WAPE percent','Revised WAPE percent'],rows,[3.3,1.8,1.8])
    para('Both columns use exactly the same 8,170 hours and inherited signed-literal-from terminals. WAPE is 100 times total absolute error divided by total absolute observed flow. The previous Central East value of 45.04% uses 8,494 hours; its matched-cohort value is 43.50%, which is the correct comparison for the revised 13.70%. Monthly, post-Dover, other terminal conventions and all lost/new masks are saved.')
    para('None of the 8,170 qualified hours passes the joint V/Q/generator-capability gate. Voltage spans 0.841226–1.120005 pu. Model AC loss averages 339.503 MW and reaches 1,024.652 MW. Assigned loss fractions reach 95.45% in zone F and 94.52% in zone E; 3,131 qualified hours have at least one zone above 25%. Those extreme values reflect unresolved network/operating assumptions, not measured NYISO losses. Parent 224 alone dissipates 0.786–172.527 model MW, averaging 39.136 MW.')
    para('Independent source accounting closes served-plus-loss load within 2.43 × 10⁻¹⁰ model MW and regional net imports within 9.07 × 10⁻¹⁰ MW. Native P changes stay below 1.86 × 10⁻¹⁰ MW. These numerical closures validate the specified equations, not the physical truth of their inputs. A 14-hour four-cell diagnostic separates impedance and allocation changes on 11 common successes; it is not an annual causal or WAPE decomposition.')
    refs('Final annual loss receipt','Final annual comparisons','Final independent physics')

    page('Conductor physics and current conventions')
    para('Thermal screening uses original same-voltage overhead conductor sections with an explicit conductor catalog, route/equivalent length, number of parallel circuits and number of subconductors per phase. The screening whitelist is a synthetic-study eligibility rule; it does not certify installed equipment. Cables, transformers, PARs, mixed voltage-changing equivalents and reduced matrix couplings need different physical models.')
    equation([sub('R','eq'),'(T) = ',frac([sub('r','sub'),'(T) L'],[sub('n','c'),sub('n','b')])], 'Req(T) = rsub(T) * L / (nc * nb)')
    equation([sub('I','sub'),' = ',frac([sub('I','series')],[sub('n','c'),sub('n','b')])], 'Isub = Iseries / (nc * nb)')
    equation(['C ',frac(['dT'],['dt']),' = ',sup(sub('I','sub'),'2'),' r(T) + ',sub('q','s'),' − ',sub('q','c'),' − ',sub('q','r')], 'C dT/dt = Isub^2 r(T) + qs - qc - qr')
    para('The heat balance is per unit length. Use the catalog’s temperature-dependent AC resistance, not a second power-scale conversion. The thermal kernel calculates convection, radiation and absorbed solar heat under declared surface and atmospheric assumptions. Steady ampacity is evaluated at a stated conductor temperature cap; a negative net cooling margin means no valid steady ampacity under those conditions.')
    para('Series current excludes lumped pi charging and is used for the modeled Joule-loss identity. From/to terminal currents include charging and use their endpoint voltage bases. Section-end current envelopes and original-terminal hardware limits are separate checks. Material segments within one modeled series element share its phase current but can have different subconductor currents and ampacity. Successive pi sections use their own recovered currents because charging can change series current along the corridor; the limiting thermal segment governs.')
    para('The existing whole-system DLR snapshots use reference R at 75°C for ACSR or 95°C for ACSS and steady ampacity at a declared cap. They do not solve a time-dependent thermal trajectory or electrothermal fixed point. Complete IEEE 738 compliance, actual span sag/clearance, route weather, surface condition, mutual effects and installed terminal limits remain outside the qualified claims.')
    refs('Engineering AC and DLR','Thermal segments','Conductor catalog','Current physical exceptions')

    page('Run a transparent DLR screening experiment')
    para('The handoff helper screens the current loss-consistent PF against steady ampacity under explicit weather. It does not change dispatch, resistance, voltage controls, or the loss-accounting equations. Its result must carry the input PF and physical-operating flags. An overload ratio is model exposure under those assumptions, not a measured violation or a mitigation benefit.')
    code("weather = table(25,0.6096,90,1000,101325, ...\n  'VariableNames',{'ambient_c','wind_m_s','wind_angle_deg', ...\n                   'solar_w_m2','pressure_pa'});\nscreen = screen_nygrid_2025_loss_pf_thermal(8508,weather);")
    para('This uniform weather is a synthetic counterfactual: 25°C air, 0.6096 m/s wind at 90 degrees to the conductor, 1,000 W/m² irradiance and 101,325 Pa pressure. For spatially varying weather, bind each value to the exact conductor-segment key using the helper’s documented contract. Never infer that the same weather occurred at all modeled routes.')
    para('The eligibility ledger excludes parent 224, source-conflicted parents 300/309/310/386, non-overhead classes, voltage-changing paths and unsupported equipment. All exclusions remain visible. The helper checks electrical-loss versus conductor-Joule consistency and reports the current/ampacity ratios at the study’s actual AC scale. It cannot certify the unresolved interiors of regional equivalents.')
    heading('The older bounded optimizer')
    code("% Explicitly historical engineering-case workflow\n[result, audit] = run_nygrid_2025_ac_dlr(8508,weather, ...\n                                      struct('solver','MIPS'));")
    para('This original API constructs npcc_nygrid_2025_ac and applies constrained ACOPF. It does not apply the latest loss revision or its filtered conductor interpretation. Archived qualification is four of eight attempted weather snapshots, with the recorded solver/environment. A different solver or machine requires inspecting the new result. Its objective is a deviation penalty, not money; current old-case results are retained for method comparison, not promoted as validated latest-model DLR benefits.')
    refs('Engineering AC and DLR','Loss revision')

    page('Design the next DLR research study')
    para('Define the research question before choosing a solver. Fixed-dispatch screening compares rating exposure at one supplied electrical state. A redispatch study must explicitly optimize allowed generation or controls and verify resulting AC feasibility. A transient study must carry conductor temperature and weather through time. A historical asset study additionally needs identified conductors, facility ceilings and route/critical-span weather.')
    tab(['Decision','Record before the experiment'],[
        ['Electrical baseline','Model snapshot, source-hour mask, dated topology, transfer policy, P/Q controls and which operating gates are required.'],
        ['Candidate assets','Original parent/section keys, conductor evidence, excluded classes and applicable terminal equipment ceilings.'],
        ['Weather and thermal state','Observation or scenario, timestamp support, route/segment mapping, wind angle, thermal initialization and gap policy.'],
        ['Perturbation','Which quantities may change; R(T), dispatch, bounds and controllers must be explicit. Preserve a control case.'],
        ['Comparison','Common hours and assets, failed-hour counts, denominators and whether the result is exposure, feasibility, curtailment or economics.'],
        ['Validation','Branch KCL, current/Joule identities, thermal residual or time-step checks, physical bounds and independent baseline reproduction.']], [1.4,5.5])
    para('Keep load and weather on a declared common calendar basis and expose both source years. Timezone and daylight-saving handling must be explicit. Do not bridge the 109 source-excluded hours or solver failures with invented thermal state. Define whether the next valid interval is reinitialized, bounded conservatively, or omitted, and report its impact.')
    para('If temperature changes R, rebuild the detailed case and its compact operators before the next electrical solve. Reconcile the coupled loss policy with those changed parameters. A joint loss-consistent bounded ACOPF/electrothermal controller is not delivered by this handoff; the one-way screening adapter must not be mistaken for it.')
    para('Priority physical gaps are the unresolved aggregate resistances, large E/F assigned losses, reactive capability/voltage controls, exact corridor equipment and complete physical flowgate mappings. Improve them using independent engineering/source evidence. Do not select impedance, losses, signs or interface membership merely because a public error metric decreases.')
    refs('P34 qualifications','Physical flowgates','Loss revision')

    page('Reproduce and inspect the evidence')
    para('The runtime package supports the included entry points and smoke examples. The complete research package adds historical construction, source receipts, saved annual profiles, controlled experiments, maps, ledgers and validation. Both annual and re-scoring commands on this page require the complete research edition. Read the inventory for exact contents; MATLAB itself and unrelated dependency caches are not distributed. Keep third-party notices with redistributed code and follow each source’s terms.')
    heading('Annual replay')
    code("addpath(fullfile(pwd,'scripts','nygrid_2025_ac_loss_revision', ...\n                 'validation'));\nreceipt = run_annual_loss_policy('released_equivalent', ...\n    'boundary_consistent','my_new_replay');")
    para('Use a new run label. The annual driver refuses to overwrite an existing attempt. It preserves all clock rows and failures, checks source/code hashes before and after the run, and writes monthly MAT arrays plus quality tables and iteration traces. The delivered 719-second timing is a historical observation on the recorded host, not a runtime guarantee.')
    heading('Post hoc comparison')
    code("python scripts/nygrid_2025_dlr_handoff/portable_scoring.py\n    --output-dir runs/comparison_2025")
    para('Enter the displayed Python command on one line from the complete research package root. This re-scores the delivered year into a new directory. The wrapper verifies the unchanged original scorer and substitutes two historical Git byte lookups with pinned package hashes. The original comparison CLI alone is not portable without Git history. The smaller runtime package omits the large annual reference arrays required for this check. A newly generated annual replay needs its own declared input/output contract rather than silently replacing the pinned reproduction inputs.')
    para('The portable scorer requires Python with NumPy, SciPy and pandas; its contract records the tested environment. Rebuilding raw public inputs may also require network services and optional tools. Frozen input/model use is offline. Old Git-dependent finalizers remain historical methods; they are not advertised as standalone commands after extraction.')
    para('Read receipts by purpose: construction and provenance; source coverage and clock; solver outcome; physical operating checks; compact/full equivalence; thermal checks; public comparisons. A receipt marked passed certifies its stated scope. A model may pass equation/recovery tests while failing the physical operating gate. Preliminary failures, superseded methods and download failures are retained as historical evidence and must not be counted as current successes.')
    refs('Final annual loss receipt','Final annual comparisons','Final hour masks','Final independent physics')

    page('Earlier structural and regional milestones')
    para('The complete archive retains the earlier local NPCC and PERFORM pathway. It is a separate lineage from the Cornell NYgrid source. Its milestones explain prior methods and decisions without supplying physical certification for the latest model. The machine-readable modification inventory includes exact source references and qualification records.')
    tab(['Milestone','Saved structure','Meaning of the result'],[
        ['S7 structural parent','143 buses, 246 branches, 62 generators','Preserved local full-NPCC structural provenance.'],
        ['S13 Phase 1A','147 buses, 253 branches, 62 generators','Four added source terminals and seven direct branches.'],
        ['S13 Phase 1B','149 buses, 262 branches, 62 generators','East Fishkill and Ladentown with nine further source circuits.'],
        ['S13 Phase 1C','170 buses, 315 branches, 62 generators','Twenty-one additional terminals and 53 source elements, including transformers.'],
        ['External reduction diagnostic','86 buses, 182 branches','Same-state mathematical identity; six bounded operating reconstructions failed.'],
        ['S11 and S12','49-bus diagnostic; 317-bus oracle','S12 has 3,531 branches and 615 generator records, including dense equivalents.'],
        ['S14 planning role','Promoted NYISO operating model pending','An S14-labelled reduction diagnostic is not promotion of this planned DLR model.']], [1.65,1.85,3.4])
    para('Do not use S12 dense equivalents as physical construction inputs. The Phase 1C implementation passed 130 implementation checks, but that numerical evidence did not qualify its bounded operating states. The original structural parent, construction overlays, diagnostic reductions and operating candidates remain separate objects.')
    heading('Regional reference packages')
    para('Package A establishes one assumed historical 2019 PERFORM reference with declared boundary and reactive-support assumptions. Its immutable source has 1,576 buses, 2,371 branches and 615 generator records. Package B is a distinct bounded historical candidate with 878 buses, 1,375 branch records and 595 generators: 855 source buses in zones D–K and 23 retained NPCC buses in A–C. Its accepted route uses MIPS with default initialization and fresh PF verification. Neither package establishes contemporary NYISO coverage or physical DLR readiness.')
    refs('Earlier project hierarchy','Phase 1C implementation','Package A foundation','Package B candidate')

    page('Earlier compact studies and paper reproduction')
    para('The smaller NY-only studies explored a different boundary and calibration design. Their scale factors, fixed boundary channels, generation policies and thermal qualifications are retained with each result. They must not be combined numerically with the present 148-bus network without an explicit transformation.')
    tab(['Study','Structure','Preserved qualification'],[
        ['Initial NY-only testbed','51 buses, 92 records, 87 active, 35 generators','Retains 46 original NY buses with a few explicit corridor terminals.'],
        ['Calibrated 2025 study','66 buses, 125 records, 112 active, 37 generators','Four interface-calibration snapshots and separate unused-hour prediction errors; 23 assumed overhead realizations.'],
        ['Source generation alternative','66 buses','14 source-qualified electrical snapshots pass; two are excluded for boundary gaps. Generation remains estimated.'],
        ['Partial Smart Path Connect','71 buses, 135 records, 117 active, 37 generators','15 electrical/sensitivity and four thermal cases pass; two electrical hours use strict IPOPT fallback. Twenty-five assumed overhead realizations.'],
        ['Later compact 2019 comparator','51 buses, 94 records, 89 active','Separate 2019 comparison variant; distinct from the earlier 92-record testbed.']], [1.75,1.8,3.35])
    para('The 66-bus calibration fitted selected interface proxies. Its small training errors do not establish unseen-hour accuracy. Independent generation and partial-SPC studies retain much larger interface prediction errors and a fixed year-end topology applied to earlier conditions. Those counterfactual electrical/thermal results remain useful methodological evidence, not historical asset validation. Optional IPOPT binaries are not supplied with this handoff.')
    heading('Reproducing the paper and repairing the measurement')
    para('The Cornell NYgrid release provides the later 2019 source pathway: 140 original buses, 227 downloaded branch records and a released 57-bus final DC model with 94 branches, 271 generators and four controllable-link records. The paper’s 233-branch description differs from the pinned download. The reconstruction records that discrepancy rather than inventing missing branches.')
    para('Five of seven plotting expressions referenced wrong branch indices. The measurement repair corrects interface extraction while preserving the solved network. The subsequent full 140-bus reconstruction, annual load verification, western diagnostics and compact comparison are included with their own receipts. The authors’ 2019 prepared local-hour labels are not interchangeable with the later 2025 physical UTC/DST clock.')
    refs('51 bus testbed','66 bus calibration','Independent generation study','71 bus partial SPC','2019 measurement repair','2019 comparator')

    page('File map and extension contract')
    tab(['Location','What the colleague should use it for'],[
        ['System Matpower Format','Public dated DC, detailed AC, compact AC and loss-consistent PF constructors/wrappers.'],
        ['scripts/nygrid_2025_dlr_handoff/runtime','Portable setup, smoke test and filtered one-way thermal screening.'],
        ['scripts/nygrid_2025_ac_loss_revision','Current loss equations, allocation policy, guards, annual replay and scorer.'],
        ['scripts/nygrid_2025_compact_ac/core','Complex reduction, voltage/current recovery, nonlinear kernel and stale-operator fingerprints.'],
        ['scripts/nygrid_2025_ac_system','Engineering branch construction, sections, generator assumptions, thermal kernels and legacy current constraints.'],
        ['output/nygrid_2025_ac_loss_revision','Final current annual results, source-equivalent review, policy evidence, masks and comparisons.'],
        ['output/nygrid_2025_ac_system','Engineering overlays, catalogs, branch metadata and archived bounded DLR results.'],
        ['output/nygrid_2025_dlr_handoff','Package inventory, portable execution proof, eligibility/exclusion ledger and final handoff validation.'],
        ['Earlier nygrid output namespaces','Frozen source-data, topology, dispatch, controls and comparison stages, identified by their own receipts.']], [2.85,4.05])
    para('Preserve frozen files when extending the study. Put revised inputs, policies, derived operators and results in a new namespace or scenario label. Carry the original parent/section keys and generation identities into every export. If a profile, parameter, device status or bus role changes, rebuild the dependent operators and update the experiment contract before comparing scores.')
    para('For an implementation handoff, retain the package hash, source model commit, MATLAB/MATPOWER version, exact invocation, weather source and units, eligibility mask, failed attempts, and the independent checks supporting the claim. These records are necessary to distinguish a reproducible synthetic DLR result from an unverified operating-system assertion.')
    heading('Research questions still open')
    para('The package supports transparent investigation of rating sensitivity and numerical model behavior. It does not establish historical thermal trajectories, actual NYISO line currents, asset-certified emergency limits, N−1 security, ramp/reserve adequacy, storage energy feasibility, or economic benefits. Those claims require their corresponding inputs and validation. The present limitations are preserved so that subsequent improvements can be attributed and audited.')
    refs('Engineering AC and DLR','Compact AC','Loss revision')

    page('Source and evidence index')
    para('All paths below are relative to the extracted package root. The original reports contain the detailed primary-source citations and receipt chains. This index identifies the relevant research stage; the machine-readable package inventory enumerates individual files and hashes.')
    for label,path in SOURCES.items():
        p=d.add_paragraph();p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_together=True
        hyperlink(p,label,path);p.add_run('\n'+path).font.size=Pt(8.5)
    md.extend(['| Evidence | Relative path |','| --- | --- |'])
    md.extend('| '+label+' | ['+path+']('+path+') |' for label,path in SOURCES.items());md.append('')
    for path in SOURCES.values():assert (ROOT/path).is_file(),path
    # Some default Word templates carry a colored title border. The handbook
    # deliberately uses plain black headings with no decorative title rule.
    for element in [d.element,d.styles.element,s.header._element,s.footer._element]:
        for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    d.save(DOCX);TEXT.write_text('\n'.join(md)+'\n',encoding='utf-8')
    receipt=dict(title=TITLE,baseline_model_commit=BASELINE,prepared_date='2026-10-13',
        chapter_count=len(pages),chapters=pages,document=str(DOCX.relative_to(ROOT)),
        markdown=str(TEXT.relative_to(ROOT)),source_links=SOURCES,native_Word_equations=True,
        no_claim_of_full_operating_or_as_built_qualification=True,
        docx_sha256=hashlib.sha256(DOCX.read_bytes()).hexdigest(),markdown_sha256=hashlib.sha256(TEXT.read_bytes()).hexdigest())
    (OUT/'handbook_build_receipt.json').write_text(json.dumps(receipt,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'docx':str(DOCX),'chapters':len(pages),'native_equations':len(d.element.xpath('.//m:oMath'))}))


if __name__=='__main__':main()
