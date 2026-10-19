"""Audit historical PERFORM lineage into current NYgrid models, without observations.

Electrical row recovery, physical device identity, dated topology, and official
flowgate membership are separate assertions. This audit establishes only the
first and records evidence for the second; it never infers contemporary gate
membership from names or assigns zero to an absent physical circuit.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from scipy.io import loadmat
from scipy.linalg import null_space

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "output/nygrid_2025_flowgate_definitions/local_model"
PINS = {
    "PERFORM/On Peak 2019 v23_Perform_NY/On Peak 2019 v23/nyiso_On_Peak_v23_shunts_as_z_load.m": "36e86cea63f1dea0872d1c2ec57f09e20d5c5ce99dad20d9977a1b014b20edb5",
    "PERFORM/Auxilliary_Perform_NY/Branch_changes_2019_to_2030.xlsx": "31d519c74f75a06a77ed11fabf535e4579fbc5c94cbb7e0a6be14cec38a054b8",
    "output/nygrid_2025_downstate/network/network_cases.mat": "89af6622ece7539d4056e3840ebda6059afcc830f0afbbdab749c43d90a503e4",
    "output/nygrid_2025_slim/network/network_cases.mat": "3c990e658464037ad36ab427150ba3c358624c1d0e3d795adb90aca91abdbc58",
    "output/nygrid_2025_downstate/network/branch_register.csv": "6a574d53b40eb2bf08dba851757a5d56ab6b4d89377b2a33e5d6319d5cd2a18f",
    "output/nygrid_2025_downstate/network/bus_register.csv": "bfab44648f3d2052e79a2a8db7c7e7390874f6bace1c575f6c94868206127f27",
    "output/nygrid_2025_downstate/network/source_bus_map.csv": "56f9bad150c960236c0982a9076d681423f5558aad47c672ab432eaffc188c72",
    "output/nygrid_2025_slim/network/branch_register.csv": "547c40ce859c358ded212f1755b879ee508a2d56c1354e7d40782b11e0bfb82d",
    "output/ny_only_package_a/source_inventory/perform_ny_branch_inventory.csv": "09898f0857ca417c99ad19516578b9b7b15adef565529fe066ae7f5df6ecd97b",
    "output/ny_only_package_a/source_inventory/perform_ny_bus_inventory.csv": "1389fef08bd6a75db5aabd7958a3a735b0dd4483181eb1c93ae31ed2fffd405f",
}
VARIANTS = ["refined_2019", "refined_2025_pre_may", "refined_2025_may21", "refined_2025_may30", "refined_2025"]


def require(value, text):
    if not bool(value):
        raise ValueError(text)


def sha(path):
    with Path(path).open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def maxabs(value):
    return float(np.max(np.abs(value), initial=0.))


def parse_matrix(text, key):
    match = re.search(r"mpc\." + key + r"\s*=\s*\[(.*?)\];", text, re.S)
    require(match is not None, "Missing source matrix " + key)
    lines = [line.split("%")[0].strip().rstrip(";") for line in match[1].splitlines()]
    return np.array([list(map(float, line.split())) for line in lines if line])


def branch_law(case):
    """Independent branch law, model MW and radians, including status/tap/phase."""
    buses, branch = np.asarray(case["bus"]), np.asarray(case["branch"])
    lookup = {int(b): j for j, b in enumerate(buses[:, 0])}
    incidence = np.zeros((len(branch), len(buses)))
    for j, row in enumerate(branch):
        incidence[j, lookup[int(row[0])]] += 1
        incidence[j, lookup[int(row[1])]] -= 1
    tap = np.where(branch[:, 8] == 0, 1., branch[:, 8])
    weight = float(case["baseMVA"]) * branch[:, 10] / (branch[:, 3] * tap)
    return weight[:, None] * incidence, -weight * np.deg2rad(branch[:, 9]), incidence


def identify_linear_target(measurement, target, constraints=None, constraint_rhs=None, tolerance=1e-10):
    """Source-constraint identifiability; no fitted data or observations permitted.

    y=Mx and Ax=b imply c*x=w*y+v*b exactly iff c lies in row([M;A]).
    Returned null perturbation is a counterexample when this fails. Numerical
    rank tolerances must be evaluated for the actual dimensional source model.
    """
    M, c = np.atleast_2d(np.asarray(measurement, float)), np.asarray(target, float).reshape(-1)
    require(M.shape[1] == len(c), "Target dimension differs from measurement")
    A = np.empty((0, len(c))) if constraints is None else np.atleast_2d(np.asarray(constraints, float))
    b = np.zeros(len(A)) if constraint_rhs is None else np.asarray(constraint_rhs, float).reshape(-1)
    require(A.shape[1] == len(c) and len(b) == len(A), "Constraint axes differ")
    C = np.r_[M, A]
    require(np.isfinite(C).all() and np.isfinite(c).all() and np.isfinite(b).all(), "Nonfinite constraint")
    N = null_space(C, rcond=tolerance)
    unresolved = c @ N
    residual = maxabs(unresolved)
    identifiable = residual <= tolerance * max(1., float(np.linalg.norm(c)))
    coeff = np.linalg.lstsq(C.T, c, rcond=tolerance)[0]
    result = dict(identifiable=bool(identifiable), nullity=int(N.shape[1]), nullspace_target_residual=residual,
                  rowspace_reconstruction_error=maxabs(coeff @ C-c), tolerance=tolerance)
    if identifiable:
        result.update(measurement_weights=coeff[:len(M)].tolist(), constraint_weights=coeff[len(M):].tolist(),
                      affine_offset=float(coeff[len(M):] @ b))
    else:
        witness = N[:, int(np.argmax(np.abs(unresolved)))]
        result.update(unobservable_perturbation=witness.tolist(), measurement_and_constraint_change=maxabs(C @ witness),
                      target_change=float(c @ witness))
    return result


def workbook_inventory(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    records, meta = [], []
    for sheet in workbook:
        rows = list(sheet.values)
        header_index = next(j for j, r in enumerate(rows) if r[0] == "ID")
        headers = rows[header_index]
        meta.append(dict(sheet=sheet.title, header_excel_row=header_index+1,
                         populated_records=sum(r[0] is not None for r in rows[header_index+1:]),
                         max_row=sheet.max_row, max_column=sheet.max_column,
                         evidence_vintage="historical_2019_or_planning_2030_not_2025_operating_topology"))
        if sheet.title not in ["2019 On Peak", "2030 On Peak"]:
            continue
        for rownum, values in enumerate(rows[header_index+1:], header_index+2):
            if values[0] is None:
                continue
            data = dict(zip(headers, values))
            records.append(dict(sheet=sheet.title, excel_row=rownum, workbook_id=str(data["ID"]),
                                source_from_bus=int(data["From Number"]), source_to_bus=int(data["To Number"]),
                                from_name=data["From Name"], to_name=data["To Name"],
                                from_kv=data["From Nom kV"], to_kv=data["To Nom kV"],
                                circuit_id=str(data["Circuit"]), status=data["Status"],
                                r_pu=data["R"], x_pu=data["X"], b_pu=data["B"], tap=data["Tap Ratio"],
                                shift_deg=data["Phase (Deg)"], rate_a=data["Lim MVA A"], rate_b=data["Lim MVA B"], rate_c=data["Lim MVA C"],
                                evidence_role="historical_or_planning_identifier_not_2025_confirmation"))
    workbook.close()
    return pd.DataFrame(records), pd.DataFrame(meta)


def build(output=OUT):
    output = Path(output)
    require(not (output / "source_crosswalk_receipt.json").exists(), "Choose a fresh output directory")
    before = {p: sha(ROOT / p) for p in PINS}
    require(before == PINS, "Pinned historical/model input changed")
    source_path = next(p for p in PINS if p.endswith("shunts_as_z_load.m"))
    source_text = (ROOT / source_path).read_text(encoding="utf-8")
    source_branch, source_bus = parse_matrix(source_text, "branch"), parse_matrix(source_text, "bus")
    branches = pd.read_csv(ROOT / "output/ny_only_package_a/source_inventory/perform_ny_branch_inventory.csv", keep_default_na=False)
    buses = pd.read_csv(ROOT / "output/ny_only_package_a/source_inventory/perform_ny_bus_inventory.csv", keep_default_na=False).set_index("source_bus")
    fullreg = pd.read_csv(ROOT / "output/nygrid_2025_downstate/network/branch_register.csv", keep_default_na=False)
    slimreg = pd.read_csv(ROOT / "output/nygrid_2025_slim/network/branch_register.csv", keep_default_na=False)
    terminalmap = pd.read_csv(ROOT / "output/nygrid_2025_downstate/network/source_bus_map.csv", keep_default_na=False).set_index("source_bus")
    full = loadmat(ROOT / "output/nygrid_2025_downstate/network/network_cases.mat", simplify_cells=True)
    slim = loadmat(ROOT / "output/nygrid_2025_slim/network/network_cases.mat", simplify_cells=True)
    require(source_branch.shape[0] == len(branches) == 2371 and len(buses) == len(source_bus) == 1576, "Historical source count changed")
    require(np.array_equal(branches.source_branch_row, np.arange(1, 2372)), "Inventory row axes changed")
    source_cols = ["source_from_bus", "source_to_bus", "source_r_pu", "source_x_pu", "source_b_pu", "source_rate_a_mva", "source_rate_b_mva", "source_rate_c_mva", "source_tap", "source_shift_deg", "canonical_status"]
    require(maxabs(branches[source_cols].to_numpy(float)-source_branch[:, :11]) < 1e-11, "Canonical inventory values differ")
    embedded = fullreg[fullreg.source_branch_row > 0].set_index("source_branch_row")
    direct = slimreg[slimreg.source_PERFORM_row > 0].set_index("source_PERFORM_row")
    require(embedded.index.is_unique and direct.index.is_unique and len(embedded) == 97 and len(direct) == 14, "Source embedding count/identity changed")
    work, workbookmeta = workbook_inventory(ROOT / "PERFORM/Auxilliary_Perform_NY/Branch_changes_2019_to_2030.xlsx")
    workbook_by_pair = {}
    for row in work.itertuples(index=False):
        key = (row.sheet, tuple(sorted((row.source_from_bus, row.source_to_bus))))
        workbook_by_pair.setdefault(key, []).append(row)
    ledger = []
    for row in branches.itertuples(index=False):
        sid, a, b = int(row.source_branch_row), int(row.source_from_bus), int(row.source_to_bus)
        mapped, retained = sid in embedded.index, sid in direct.index
        modelrow = int(embedded.loc[sid, "model_branch_row"]) if mapped else 0
        term_a = int(terminalmap.loc[a, "model_bus"]) if a in terminalmap.index else 0
        term_b = int(terminalmap.loc[b, "model_bus"]) if b in terminalmap.index else 0
        if mapped:
            frow = full["refined_2019_case"]["branch"][modelrow-1]
            require(np.array_equal(frow[:2], [term_a, term_b]), "Source/model endpoint direction changed")
            require(maxabs(frow[2:11]-source_branch[sid-1, 2:11]) < 1e-11, "Embedded source electrical parameters changed")
        candidates = {year: workbook_by_pair.get((year+" On Peak", tuple(sorted((a,b)))), []) for year in ["2019", "2030"]}
        ledger.append(dict(source_branch_row=sid, raw_device_key=row.device_key, raw_circuit_id=row.raw_circuit_id,
                           raw_identity_candidates=row.raw_identity_candidates, individual_RAW_identity_resolved=int(row.individual_circuit_identity_resolved),
                           RAW_identity_basis=row.identity_confidence, source_from_bus=a, source_to_bus=b,
                           source_from_name=buses.loc[a, "source_bus_name"], source_to_name=buses.loc[b, "source_bus_name"],
                           source_from_kv=float(buses.loc[a,"base_kv"]), source_to_kv=float(buses.loc[b,"base_kv"]),
                           source_from_zone=buses.loc[a,"source_zone"], source_to_zone=buses.loc[b,"source_zone"],
                           r_pu=row.source_r_pu,x_pu=row.source_x_pu,tap=row.source_tap,shift_degree=row.source_shift_deg,
                           source_status=int(row.canonical_status),full191_branch_row=modelrow,
                           slim148_direct_branch_row=int(direct.loc[sid,"model_branch_row"]) if retained else 0,
                           source_embedded_from_model_bus=term_a,source_embedded_to_model_bus=term_b,
                           model_recovery=("direct_retained_historical_row" if retained else "affine_recoverable_historical_row") if mapped else "not_embedded_no_physical_flow_estimate",
                           source_to_full_orientation=1 if mapped else "",recovery_sign_from_terminal=1 if mapped else "",recovery_sign_to_terminal=-1 if mapped else "",
                           recovery_meter_end_scope="lossless_DC_terminal_flows_opposite_no_AC_loss_reconstruction" if mapped else "unresolved",
                           source_vintage="PERFORM_NY_2019_v23_RAW_export_2022_09_19",
                           physical_2025_identity_confirmed=False,physical_2025_membership_confirmed=False,
                           physical_2025_terminal_alias_basis="separate_dated_primary_evidence_required",
                           workbook_2019_endpoint_candidates=";".join(x.workbook_id for x in candidates["2019"]),
                           workbook_2030_endpoint_candidates=";".join(x.workbook_id for x in candidates["2030"]),
                           workbook_candidates_are_identity_proof=False,
                           absence_policy="unmapped_is_unknown_never_zero_or_aggregate_substitute",
                           **{v+"_status":int(embedded.loc[sid,v+"_status"]) if mapped else "" for v in VARIANTS}))
    ledger = pd.DataFrame(ledger)
    checks = []
    mappedrows = embedded.model_branch_row.to_numpy(int)-1
    for variant in VARIANTS:
        case, recovery = full[variant+"_case"], slim["recovery"][variant]
        F, phase, incidence = branch_law(case)
        L, d = recovery["angle_lift"], recovery["angle_offset"]
        independent_H, independent_h = F @ L, F @ d + phase
        operator_error = maxabs(independent_H-recovery["branch_angle_matrix"])
        offset_error = maxabs(independent_h-recovery["branch_flow_offset"])
        interior = np.asarray(recovery["eliminated_full_indices_1based"], int)-1
        interior_kcl_operator_error = maxabs(incidence[:,interior].T @ independent_H)
        interior_kcl_offset_error = maxabs(incidence[:,interior].T @ independent_h)
        require(operator_error < 1e-8 and offset_error < 1e-8 and interior_kcl_operator_error < 1e-8 and interior_kcl_offset_error < 1e-8, "Affine recovery failed")
        # PAR211 is retained, so removing it does not change eliminated-interior equations.
        par_case = dict(case, branch=np.array(case["branch"], copy=True))
        require(np.array_equal(par_case["branch"][210,:2], [48,100]), "PAR representation row changed")
        par_case["branch"][210,10] = 0
        Fpar, ppar, _ = branch_law(par_case)
        par_error = maxabs((Fpar @ L)[mappedrows]-independent_H[mappedrows])
        require(par_error == 0, "Mapped circuit recovery changed by retained PAR substitution")
        checks.append(dict(variant=variant,recoverable_historical_rows=len(mappedrows),all410_operator_error_MW_per_rad=operator_error,
                           all410_offset_error_model_MW=offset_error,eliminated_KCL_operator_error_MW_per_rad=interior_kcl_operator_error,
                           eliminated_KCL_offset_error_model_MW=interior_kcl_offset_error,PAR211_substitution_mapped97_error_MW_per_rad=par_error,
                           units="radians_in_model_MW_out",physical_2025_validation=False))
    fixtures = {
        "parallel_bundle_sum_identifiable":identify_linear_target([[1,1]],[1,1]),
        "parallel_individual_unidentifiable":identify_linear_target([[1,1]],[1,0]),
        "individual_with_source_known_equal_sharing_constraint":identify_linear_target([[1,1]],[1,0],[[1,-1]],[0]),
        "serial_difference_known_side_injection":identify_linear_target([[1,0]],[0,1],[[1,-1]],[5]),
    }
    require(fixtures["parallel_bundle_sum_identifiable"]["identifiable"], "Sum fixture failed")
    require(not fixtures["parallel_individual_unidentifiable"]["identifiable"], "Individual unresolved fixture failed")
    require(abs(fixtures["serial_difference_known_side_injection"]["affine_offset"]+5) < 1e-10, "Affine fixture failed")
    roles = fullreg.copy()
    roles["physical_2025_gate_member_inferred"] = False
    roles["physical_identity_policy"] = np.where(roles.source_branch_row>0,"historical_PERFORM_row_only_separate_2025_identity_required","aggregate_or_assumption_no_individual_physical_identity")
    roles.loc[roles.model_branch_row==211,"physical_identity_policy"] = "operational_PAR_substitution_original_AC_row_inactive_two_ideal_PS33_PS34_transfers_not_a_physical_AC_line"
    output.mkdir(parents=True,exist_ok=True)
    ledger.to_csv(output/"PERFORM_source_to_model.csv",index=False)
    roles.to_csv(output/"full191_physical_identity_boundaries.csv",index=False)
    work.to_csv(output/"historical_2019_planning_2030_workbook_rows.csv",index=False)
    workbookmeta.to_csv(output/"historical_workbook_sheet_inventory.csv",index=False)
    pd.DataFrame(checks).to_csv(output/"affine_recovery_checks.csv",index=False)
    (output/"identifiability_fixtures.json").write_text(json.dumps(fixtures,indent=2)+"\n")
    after = {p: sha(ROOT/p) for p in PINS}
    require(after==before,"Protected source changed during audit")
    receipt=dict(passed=True,created_utc=datetime.now(timezone.utc).isoformat(),scope="historical_source_to_current_model_lineage_no_2025_flowgate_promotion",
                 source_rows=len(ledger),full191_exact_historical_rows=int((ledger.full191_branch_row>0).sum()),
                 slim148_direct_historical_rows=int((ledger.slim148_direct_branch_row>0).sum()),
                 recovered_historical_rows=int((ledger.model_recovery=="affine_recoverable_historical_row").sum()),
                 unmapped_historical_rows=int((ledger.full191_branch_row==0).sum()),
                 mapped_RAW_individually_resolved=int(ledger.loc[ledger.full191_branch_row>0,"individual_RAW_identity_resolved"].sum()),
                 mapped_RAW_ambiguous=int((ledger.loc[ledger.full191_branch_row>0,"individual_RAW_identity_resolved"]==0).sum()),
                 parent_aggregate_rows_without_physical_source_identity=225,physical_2025_gate_membership_established=False,
                 internal_interface_observations_read=False,observational_fit_used=False,source_hashes_unchanged=True,
                 dependencies=before,script_sha256=sha(__file__),validation=checks,
                 operator_contract={"formula":"H=c F_full L; h=c(F_full d+phase); flow_model_MW=H theta_retained_rad+h",
                   "scope":"only source-backed complete term c on full191; zero eliminated PD/QD/GS/BS/no generator/DC terminal",
                   "physical_unknowns":"2019/2030 planning source is not 2025 membership, outage, metering, line-resection or equipment data",
                   "unknown_term":"fail_closed_for_full_gate; partial_sum_is_separately_labelled_and_must_not_be_scored_as_complete_gate",
                   "aggregate_identifiability":"c annihilates null([M;A]); M,A,b require independent physical source evidence; absent such constraints do not solve by curve fitting",
                   "metering":"from/to sign is opposite in lossless DC; actual metering terminal must still be sourced; AC terminal loss differences unrecoverable"})
    receipt["outputs"]={p.name:sha(p) for p in output.iterdir() if p.is_file()}
    (output/"source_crosswalk_receipt.json").write_text(json.dumps(receipt,indent=2)+"\n")
    print(json.dumps({k:receipt[k] for k in ["passed","source_rows","full191_exact_historical_rows","slim148_direct_historical_rows","recovered_historical_rows","unmapped_historical_rows","mapped_RAW_individually_resolved","mapped_RAW_ambiguous"]},indent=2))


if __name__ == "__main__":
    parser=argparse.ArgumentParser()
    parser.add_argument("--output",type=Path,default=OUT)
    build(parser.parse_args().output)
