"""Read-only independent extraction of six matched EIA923 wind plants.

Run with the bundled spreadsheet Python. No workbook is written or modified.
"""
from datetime import datetime,timezone
import hashlib,json
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'output/nygrid_2025_postrepair_diagnosis/boundary'
BOOK=ROOT/'output/nygrid_2025_western_audit/generation/sources/EIA923_Schedules_2_3_4_5_M_12_2025_Final.xlsx'
PLANTS={56618,56619,56857,56901,56904,59629}


def sha(p):
    with p.open('rb')as f:return hashlib.file_digest(f,'sha256').hexdigest()


def main():
    before=sha(BOOK);w=load_workbook(BOOK,read_only=True,data_only=True)
    sheet=w['Page 1 Generation and Fuel Data']
    iterator=sheet.iter_rows(values_only=True);top=[next(iterator)for _ in range(6)]
    assert '2025 Final Data' in str(top[1][0])
    headers=[' '.join(str(v).split())for v in top[5]];index={v:i for i,v in enumerate(headers)}
    cols=[i for i,h in enumerate(headers)if h.startswith('Netgen ')];assert len(cols)==12
    rows=[]
    for number,values in enumerate(iterator,7):
        if values[index['Plant Id']] not in PLANTS:continue
        if values[index['Reported Fuel Type Code']]!='WND':continue
        assert values[index['Plant State']]=='NY' and values[index['Balancing Authority Code']]=='NYIS'
        assert values[index['Reported Prime Mover']]=='WT' and values[index['YEAR']]==2025
        months=[float(values[i])for i in cols]
        annual=float(values[index['Net Generation (Megawatthours)']])
        assert abs(sum(months)-annual)<1e-6
        rows.append(dict(plant_id=int(values[index['Plant Id']]),plant_name=values[index['Plant Name']],
            prime_mover='WT',fuel='WND',state='NY',balancing_authority='NYIS',source_year=2025,
            respondent_frequency=values[index['Respondent Frequency']],source_excel_row=number,
            annual_MWh=annual,monthly_MWh=months,
            annual_source_cell=f"{get_column_letter(index['Net Generation (Megawatthours)']+1)}{number}",
            monthly_source_range=f'{get_column_letter(cols[0]+1)}{number}:{get_column_letter(cols[-1]+1)}{number}'))
    assert len(rows)==6 and {r['plant_id']for r in rows}==PLANTS
    notes=[]
    for row in w['Page 7 File Layout'].iter_rows(values_only=True):
        text=' | '.join(str(v)for v in row if v is not None)
        if any(x in text.lower()for x in ['respondent','annual respondent','monthly respondent']):notes.append(text)
    w.close();assert sha(BOOK)==before
    result=dict(passed=True,created_utc=datetime.now(timezone.utc).isoformat(),workbook=BOOK.relative_to(ROOT).as_posix(),
        workbook_sha256=before,sheet='Page 1 Generation and Fuel Data',header_row=6,net_generation_unit='MWh',
        source_rows=rows,workbook_frequency_layout_notes=notes,annual_net_generation_MWh=sum(r['annual_MWh']for r in rows),
        workbook_modified=False,extraction='openpyxl read_only=True data_only=True; exact plant/fuel/state/BA/year filters')
    OUT.mkdir(parents=True,exist_ok=True);(OUT/'independent_wind_workbook_readback.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({k:result[k]for k in ['passed','workbook_sha256','annual_net_generation_MWh','workbook_frequency_layout_notes','source_rows']},indent=2))


if __name__=='__main__':main()
