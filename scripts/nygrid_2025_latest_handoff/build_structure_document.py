"""Author the current model structure only; operational instructions are separate."""
from pathlib import Path
import sys
import json
import hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT/'scripts/nygrid_2025_dlr_handoff'))
from build_handbook import table, math, sub, sup, frac, hyperlink

OUT = ROOT/'output/nygrid_2025_latest_handoff'
TITLE = 'NYgrid 2025 current model structure'
BASE = '2e0fcd5c3f0ae39d04f27a733ce30eb148b0b243'

def main():
    d = Document(); s = d.sections[0]
    s.page_width=Inches(8.5); s.page_height=Inches(11)
    s.top_margin=Inches(.67); s.bottom_margin=Inches(.67)
    s.left_margin=Inches(.8); s.right_margin=Inches(.8)
    s.header_distance=Inches(.28); s.footer_distance=Inches(.28)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        st=d.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_after=Pt(7)
    d.styles['Normal'].font.size=Pt(11); d.styles['Normal'].paragraph_format.line_spacing=1.06
    d.styles['Title'].font.size=Pt(25); d.styles['Title'].font.bold=True
    d.styles['Subtitle'].font.size=Pt(12)
    d.styles['Heading 1'].font.size=Pt(18)
    d.styles['Heading 2'].font.size=Pt(12); d.styles['Heading 2'].font.bold=True
    for name in ['Heading 1','Heading 2']:
        d.styles[name].paragraph_format.keep_with_next=True
    s.header.paragraphs[0].text='NYgrid 2025  |  Current model structure'
    s.header.paragraphs[0].style='Caption'
    f=s.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    f.add_run('13 October 2026  |  ')
    field=OxmlElement('w:fldSimple');field.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instr','PAGE');f._p.append(field)
    d.core_properties.title=TITLE
    d.core_properties.subject='Current electrical and thermal model structure for research'
    d.core_properties.author='NYgrid research project'
    md=['# '+TITLE,'','Prepared 13 October 2026. Current model snapshot '+BASE+'.','']
    pages=[]
    def p(text): d.add_paragraph(text);md.extend([text,''])
    def h(text): d.add_paragraph(text,'Heading 2');md.extend(['### '+text,''])
    def page(text):
        d.add_page_break();d.add_paragraph(text,'Heading 1');md.extend(['## '+text,'']);pages.append(text)
    def tab(headers,rows,widths):
        table(d,headers,rows,widths)
        md.extend(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |'])
        md.extend('| '+' | '.join(str(v) for v in row)+' |' for row in rows);md.append('')
    def eq(items,text):math(d,items);md.extend(['```text',text,'```',''])

    d.add_paragraph(TITLE,'Title')
    d.add_paragraph('Electrical network and conductor thermal representation','Subtitle')
    p('This document describes the model supplied for continued dynamic line rating research: a 148-bus, loss-consistent AC power-flow model with recoverable detailed branch currents and conductor realizations constrained by electrical resistance and reactance. It explains the current network, operating inputs, equations, controls and thermal interpretation. Package setup and run commands are in START_HERE.md.')
    p('The compact and detailed networks are two representations of the same current model. Four dated topology states select the network applicable to each 2025 hour. The case is a synthetic research equivalent; its electrical and thermal assumptions do not establish an observed NYISO operating state or installed asset rating.')
    tab(['Quantity','Current structure'],[
        ['Retained AC buses','148 in every dated state'],
        ['Detailed AC buses','266 before full-service Dover; 267 afterward'],
        ['Detailed AC branch rows','485 before full-service Dover; 488 afterward'],
        ['Parent corridor records','410 before full-service Dover; 413 afterward'],
        ['Generator rows','1,464, with hourly status and bounds'],
        ['Prescribed transfer rows','6 MATPOWER DC-line records'],
        ['Time axis','8,760 hours; 8,651 source-eligible hours'],
        ['Electrical base','100 MVA; bus-specific nominal kV; 60 Hz']], [2.2,4.7])
    picture=OUT/'structure/current_model_schematic.png'
    if picture.exists():
        d.add_picture(str(picture),width=Inches(6.9))
        md.extend(['![Current model representations](model/structure/current_model_schematic.png)',''])
    p('Parent corridors identify modeled facilities or equivalents. Detailed branch rows include numerical line sections and equipment records. Reduced admittance couplings are mathematical objects and have no independent conductor or thermal rating.')
    pages.append('Overview')

    page('Network and dated topology')
    p('The 148 retained buses comprise 54 New York, 36 ISONE, 31 Ontario and 27 PJM buses. They include 140 NPCC boundary and internal buses plus eight New York 345-kV station ports. The additional ports retain junction voltages without assigning them extra demand or generation solely because they are retained.')
    tab(['Retained station port','Bus ID','Retained station port','Bus ID'],[
        ['East Fishkill',10858,'Wood Street',10902],
        ['Ladentown',10774,'Dunwoodie',10593],
        ['East Garden City',10589,'Gowanus',11129],
        ['Rainey',11130,'Cricket Valley',20001]], [2.5,.95,2.5,.95])
    p('In the detailed representation, 43 additional Downstate terminals and 75 numerical section buses are passive and eliminated. The post-Dover state adds one passive regulated port. Thus 118 or 119 eliminated voltages supplement the same 148 retained voltages. Numerical section buses represent discretization, not extra substations.')
    tab(['State begins','Topology state','Detailed buses / branches'],[
        ['Start of local 2025','pre_may','266 / 485'],
        ['21 May 2025 00 EDT','may21','266 / 485'],
        ['30 May 2025 00 EDT','may30','266 / 485'],
        ['19 Dec 2025 23 EST','post_dover','267 / 488']], [2.35,1.9,2.65])
    p('The May boundaries activate the modeled Gowanus–Greenwood and Goethals–Fox Hills receiving paths. The Dover boundary is 20 December 2025 at 04 UTC. These dates are implemented switching policies. An hour selects exactly one state; the states are not separate model versions.')
    h('Detailed branch representation')
    p('The detailed network contains overhead corridors, cable and urban equivalents, transformer-compatible paths, composite voltage-changing equivalents, PAR equipment and inactive records. The section map records each parent, its numerical sections and its physical from/to terminals. Inactive rows remain in the axis and must be distinguished from active zero-flow rows.')
    p('Fixed shunts and declared charging compensation remain part of the admittance model. Parent 224 is a released Gilboa–Buchanan network equivalent with total R = 0.02 pu, X = 0.0147 pu and B = 0; it remains electrical-only under the current conductor catalog and geometry limits.')

    page('Generation and controlled transfers')
    p('The 1,464 generator rows contain 1,436 native resource rows, 27 external regional equivalents and one signed HQ equivalent. A generator row is distinct from its electrical bus ID and reporting identity. Multiple rows can share a bus. Hourly active-power schedules, capability bounds and status are supplied on a fixed row axis; the active count changes with the hour.')
    h('Active and reactive controls')
    p('Native real-power schedules and the signed HQ injection are fixed in the loss-consistent solve. ISONE, Ontario and PJM equivalents provide their assigned AC losses using fixed participation weights and finite original real-power bounds. External delivered-load proxies remain fixed. The reference angle is set at bus 50; a numerical reference source is not an unlimited provider of active-power losses.')
    p('PV buses retain their stated voltage magnitudes. The solver determines the associated reactive injections without enforcing generator Q limits through bus-type switching. The inherited capability assumptions use a 0.95 power-factor envelope and apparent-power circles where specified. The audit separately checks the 0.95–1.05 pu voltage range, reactive bounds and generator capability.')
    h('Prescribed exchanges')
    p('Six MATPOWER DC-line records prescribe active-power transfers: four external-transfer representations and two Ontario PAR-flow substitutions. They are six modeling records, not six asserted physical HVDC assets. Their scheduled signs, endpoint injections and reactive conventions are retained. Their declared real-power losses are zero in the current accounting policy.')
    tab(['Channel','Model from bus','Model to bus'],[
        ['CSC', '21 MILLSTONE PT', '80 NORTHPORT'],
        ['NEPTUNE', '124 BRANCHBURG', '79 RAV A-3'],
        ['HTP and VFT', '125 LINDON', '81 GOETHALS'],
        ['Ontario PAR 25118 and 25119', '100 PHAS', '48 MOSES E']], [2.5,2.2,2.2])
    p('These are model endpoint buses, not asserted literal asset terminals. Positive transfer is from-to; a negative target reverses it. Endpoint reactive injections are fixed at zero for these six records.')
    p('The Ontario substitutions disable the replaced aggregate AC transfer path. In the post-Dover state, two separate parallel PAR-bank records use the bounded hourly phase reconstruction. These prescribed controls condition the power flow; the solve does not optimize their schedules or infer actual tap positions.')
    h('Meaning of the input dispatch')
    p('The native dispatch is source-conditioned and capacity bounded. It is not a unit-commitment solution: ramps, reserves, startup decisions, minimum stable output and storage energy trajectories are not established. External demand and reactive behavior are synthetic assumptions. The model is suitable for controlled research scenarios only when these assumptions are retained or explicitly replaced.')

    page('AC equations and exact recovery')
    p('Let V be the complex bus-voltage vector, Y the detailed bus-admittance matrix and S the net complex injection on the stated base. The coupled solver evaluates voltage-dependent delivered demand and external loss participation together with the AC equations.')
    eq(['S(V) = diag(V) ',sup('(YV)','*')], 'S(V) = diag(V) conj(Y V)')
    p('Partition buses into retained r and eliminated e. The eliminated buses have zero independent complex injection under the frozen placement. Their charging and fixed shunts remain in Y. Eliminating those voltages gives a complex-admittance reduction rather than a DC susceptance approximation.')
    eq([sub('V','e'),' = −',sup(sub('Y','ee'),'−1'),sub('Y','er'),sub('V','r')], 'Ve = -inv(Yee) Yer Vr')
    eq([sub('Y','red'),' = ',sub('Y','rr'),' − ',sub('Y','re'),sup(sub('Y','ee'),'−1'),sub('Y','er')], 'Yred = Yrr - Yre inv(Yee) Yer')
    p('The 148 retained voltages include a fixed reference angle and controlled PV magnitudes; they are not 296 free real unknowns. Bus types depend on hourly generator admission. A linear lift L reconstructs all detailed voltages. The current operators include branch series admittance, charging, transformer taps and phase shifts, preserving full-network bus and branch ordering.')
    eq(['V = L',sub('V','r'),';   ',sub('I','f'),' = ',sub('Y','f'),'L',sub('V','r'),';   ',sub('I','t'),' = ',sub('Y','t'),'L',sub('V','r')], 'V = L Vr; If = Yf L Vr; It = Yt L Vr')
    p('A separate series-current operator removes the terminal charging contribution for conductor Joule heating. Terminal currents remain relevant to terminal equipment and the declared rating screen. Current conversion uses the nominal voltage at the appropriate endpoint and the three-phase electrical base.')
    eq([sub('I','base'),' = ',frac(['100 × ',sup('10','6')],['√3 × ',sub('V','base,kV'),' × ',sup('10','3')])], 'Ibase_A = baseMVA * 1e6 / (sqrt(3) * base_kV * 1e3)')
    p('All variable loss-dependent injections remain on retained buses. Adding a new load, generator or controller to an eliminated bus changes the reduction contract and requires rebuilding the retained set and operators. Thermal calculations use recovered detailed currents, never reduced admittance edges.')

    page('Hourly inputs and loss accounting')
    p('The model stores one annual axis of 8,760 UTC-referenced timestamps spanning local calendar year 2025. A source mask admits 8,651 hours and excludes 109. The bundle contains zonal and bus demand allocations, generator schedules and bounds, prescribed transfers, phase controls, state IDs and eligibility flags. Excluded hours are not replaced with artificial zero inputs.')
    h('Power scale and source meaning')
    p('The fixed power scale is gamma = 0.30012075175285535, based on a declared 9,561-MW model anchor divided by the 31,857.1773-MW maximum simultaneous public zonal load. Model MW equals gamma times public MW. Aggregate model MW may be divided by gamma for public-series comparison. AC amperes are never divided by gamma.')
    p('NY zonal actual load is interpreted as a generation-plus-net-interchange balance. The current model partitions its frozen scaled target into delivered demand and assigned AC losses. Delivered demand is an inferred model quantity, not an observed end-use load series. Within-zone spatial shares and reactive power factors are fixed assumptions.')
    eq([sub('P','served,z'),' + ',sub('P','loss,z'),'(V) = ',sub('P','gross,z')], 'served_PD_z + allocated_AC_loss_z(V) = frozen_gross_P58_z')
    p('Within-NY corridor losses are shared equally by endpoint zones; a corridor within one zone assigns its full loss there. External internal corridors and NY-border corridors assign their losses to the appropriate external region. HQ is fixed. These are declared accounting rules, not observed meter allocations.')
    h('Simultaneous numerical solve')
    p('The solver jointly updates retained voltages and loss-dependent injections using analytic derivatives and a domain-preserving line search. It rejects negative NY delivered demand and external participation that exceeds finite bounds. It does not clip losses, expand capability or fall back to unlimited reference-generator supply.')
    p('Successful accounting requires both the AC equations and regional/boundary identities to close. A failed line search or a delivered-load boundary is a failed attempted solve, not proof that every possible operating solution is infeasible. Failed iterates and excluded source hours must not be treated as valid solved states.')

    page('Conductor and thermal representation')
    p('A synthetic conductor realization is acceptable without identifying an installed circuit. Each accepted overhead realization specifies catalog material, resistance temperature, diameter, GMR, circuit count, bundle count, phase geometry and equivalent length. Its forward resistance and reactance must match the current electrical section. Resistance and reactance are held fixed during selection; charging mismatch is disclosed separately.')
    eq(['R = ',frac([sub('r','sub'),'(T) ℓ'],[sub('n','c'),sub('n','b')]),';   X = ',frac(['2πf × 2 × ',sup('10','−7'),' ℓ'],[sub('n','c')]),' ln(',sub('D','m'),' / ',sub('GMR','bundle'),')'], 'R_ohm = r_sub(T)*length/(n_c*n_b); X_ohm = 2*pi*f*2e-7*length/n_c*ln(GMD/GMR_bundle)')
    p('Here R and X are ohms, length is metres and r_sub is ohms per metre. Convert per-unit impedance using Z_base = V_base,kV² / baseMVA. n_c counts parallel circuits and n_b counts subconductors per phase. X/R constrains geometry; R fixes length for a chosen multiplicity. Several conductor, bundle, circuit and length combinations may reproduce one R/X pair. Catalog temperature and voltage-dependent geometry/length limits define the admissible synthetic family.')
    p('The selector also requires active same-voltage overhead sections with compatible taps, phase shifts and consistent segment metadata. Known cable conflicts at parents 300, 309, 310 and 386 are excluded. Parent 224 fails the current 345-kV catalog/geometry X/R condition; this is a conditional no-fit result, not a universal impossibility proof.')
    tab(['Runtime topology','Eligible parents','Branch rows','Material segments'],[
        ['Before full-service Dover',179,251,251],
        ['After full-service Dover',180,252,253]], [2.7,1.2,1.25,1.75])
    p('Current electrical parameters include both realizations preserving imported R/X and declared engineering constructions. The eligibility ledger records that distinction. The screen does not change either category merely to produce an acceptable thermal utilization.')
    eq([sup(sub('I','sub'),'2'),sub('r','sub'),'(T) + ',sub('q','solar'),' = ',sub('q','conv'),'(T) + ',sub('q','rad'),'(T)'], 'I_sub^2 r_sub(T) + q_solar = q_conv(T) + q_rad(T)')
    p('Explicit air temperature, wind speed, wind angle, solar irradiance and pressure drive the steady heat balance. Conductor current is the recovered series current divided by n_c n_b. Material segments sharing one lumped series element share that element current; successive electrical pi sections use their own recovered currents. The segment temperature limit determines ampacity. This is a one-way thermal screen, not an R(T)-coupled power flow or DLR optimization.')

    page('Data axes and interpretation limits')
    tab(['Object or axis','Meaning'],[
        ['148 retained bus rows','Bus IDs and retained voltages; row number is not the bus ID.'],
        ['Detailed bus and branch rows','266/267 voltages and 485/488 original-section currents for the selected state.'],
        ['1,464 generator rows','Fixed generator identity axis with corrected electrical landing and hourly capability.'],
        ['Parent and section maps','Parent ID, branch key, numerical sections and physical terminal markers.'],
        ['Thermal segment rows','Catalog, geometry, material, reference temperature and section-current membership.'],
        ['Hour index and timestamp','One of 8,760 positions, accompanied by source eligibility and topology state.']], [2.15,4.75])
    h('Separate qualification flags')
    p('result.success and audit.loss_allocation_gate establish numerical and accounting closure; audit.reduction_and_equation_gate checks recovered detailed equations. audit.VQ_and_generator_assumption_gate checks operating limits; audit.passed additionally requires those limits. Thermal eligibility establishes a synthetic conductor mapping; utilization evaluates specified weather. None of these flags alone establishes all the others.')
    p('The source eligibility mask does not certify an AC operating point. A point may satisfy the coupled equations and loss balances while violating the voltage or generator limits, or while exceeding a scenario ampacity. Preserve each qualification flag alongside its outputs rather than combining them into a single unqualified success label.')
    p('The model supports reproducible studies of its AC equations, boundary accounting, reduction and assumed conductor thermal behavior. It does not establish actual critical-span weather, installed asset identity, terminal-limited facility ratings, a feasible dispatch history, N−1 security or economic benefit from DLR. Scenario changes must retain the electrical and thermal bases and must be requalified.')
    h('Machine-readable structure')
    p('The package includes the current model bundle and its receipt, model/structure ledgers, DATA_DICTIONARY.md and an integrity manifest. The structure ledgers provide the full bus, parent, section, generator and transfer axes behind this description. START_HERE.md contains the public entry points and execution examples; this document is limited to the model itself.')
    for el in list(d.styles.element.iter())+list(d.element.iter()):
        if el.tag.endswith('}pBdr') and el.getparent() is not None: el.getparent().remove(el)
    docx=ROOT/'NYGRID_2025_LATEST_MODEL_STRUCTURE.docx'; text=ROOT/'NYGRID_2025_LATEST_MODEL_STRUCTURE.md'
    d.save(docx); text.write_text('\n'.join(md).rstrip()+'\n',encoding='utf-8')
    result={'title':TITLE,'source_model_commit':BASE,'planned_pages':pages,'scope':'current_model_structure_only',
            'files':{p.name:{'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size} for p in [docx,text]}}
    OUT.mkdir(parents=True,exist_ok=True)
    (OUT/'structure_document_receipt.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(result,indent=2))

if __name__=='__main__': main()
